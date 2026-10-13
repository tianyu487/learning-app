#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成「600 项目总表」。
- 5 学科 × 10 等级 × 12 个 = 600
- 输出:项目总表_600.xlsx(精装:分级配色 / 冻结 / 筛选 / 覆盖矩阵)+ 项目总表_600.csv
运行:  python 生成600项目.py
依赖:  openpyxl   (pip install openpyxl --break-system-packages)
"""
import csv, math

COLS = ["项目编号","项目名称","学科","阶级","主题/单元","一句话简介","核心知识点",
        "出镜人物","方法论","动手活动","产出形式","物料包清单","物料成本(元)",
        "课时(分钟)","适龄/年级","难度(1-5)","前置项目编号","自闭症适配",
        "定价(元)","制作状态","资产链接","备注"]

VARIANTS = ["入门探索","动手做一做","观察与发现","小实验","拆解原理","测量与记录",
            "对比挑战","搭建作品","讲给你听","创意改造","闯关游戏","综合复盘"]

AGE = {1:"小学一年级",2:"小学二年级",3:"小学三年级",4:"小学四年级",5:"小学五年级",
       6:"小学六年级",7:"初一",8:"初二",9:"初三/高一",10:"高二/竞赛"}

# 学科: 名称, 代码, 出镜人物, 物料, 成本基数, 10个等级主题
SUBJECTS = [
 ("数学","MA","高斯爷爷","数字卡、积木、方格纸",4,
  ["数感与计数","加减与凑十","乘除与规律","图形与对称","分数与小数",
   "面积与体积","方程与未知数","比例与函数入门","几何证明与逻辑","代数综合与奥赛"]),
 ("物理","PH","工程师阿木","电池、磁铁、小车、木块",6,
  ["推与拉(力的启蒙)","浮与沉","光与影","声音的秘密","简单机械",
   "电路初步","电与磁","力与运动","能量与功","综合实验与竞赛"]),
 ("化学","HX","居里夫人","小苏打、白醋、盐、紫甘蓝",5,
  ["认识物质","水的三态","溶解与混合","酸与碱(变色)","气体探秘(二氧化碳)",
   "元素与符号","化学反应与方程","金属与氧化","溶液与浓度","化学计量与探究"]),
 ("AI技术栈","AI","导演小创","AI工具账号、模板(轻物料)",0,
  ["认识AI伙伴","用AI画画","用AI讲故事","指令与提示词","创意与AI视频",
   "数据与图表","编程思维入门","智能体小项目","AI产品设计","AI综合作品与发布"]),
 ("文·经典","WX","李白","诗词卡、绘本、纸笔",3,
  ["古诗启蒙(跟李白)","成语故事","经典朗读","看图说话与表达","写一段小故事",
   "论语与做人","古文入门","名著选读","思辨与演讲","经典写作与创作"]),
]

def meth(L):  return "苏格拉底·游戏" if L<=3 else "费曼·项目制" if L<=6 else "项目制·刷题" if L<=8 else "刷题·费曼"
def out(L):   return "作品+讲解视频" if L<=3 else "作品+举证视频" if L<=6 else "实验记录+演讲" if L<=8 else "研究报告+展示"
def aut(L):   return "是" if L<=4 else "部分"
def diff(L):  return math.ceil(L/2)
def price(code,L): return "高阶单列" if code=="AI" and L>=7 else str(29+3*L)

def build_rows():
    rows=[]
    for name,code,person,mat,base,themes in SUBJECTS:
        for L in range(1,11):
            theme=themes[L-1]
            for nn in range(1,13):
                v=VARIANTS[nn-1]
                pid=f"{code}-{L:02d}-{nn:02d}"
                prereq="" if L==1 else f"{code}-{L-1:02d}-{nn:02d}"
                cost = 0 if code=="AI" else base+L
                rows.append([pid, f"{theme} · {v}", name, L, theme,
                             f"围绕「{theme}」的{v}", theme, person, meth(L), v, out(L),
                             mat, cost, 30+5*L, AGE[L], diff(L), prereq, aut(L),
                             price(code,L), "规划", "", ""])
    return rows

def write_csv(rows, path="项目总表_600.csv"):
    with open(path,"w",newline="",encoding="utf-8-sig") as f:
        w=csv.writer(f); w.writerow(COLS); w.writerows(rows)
    print("CSV ->",path, len(rows),"rows")

def write_xlsx(rows, path="项目总表_600.xlsx"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb=Workbook(); ws=wb.active; ws.title="项目总表"
    head_fill=PatternFill("solid",fgColor="1D9E75"); head_font=Font(bold=True,color="FFFFFF",name="Arial")
    band={1:"FCEFE1",2:"FCEFE1",3:"F6E2C8",4:"F6E2C8",5:"EAF3DE",6:"EAF3DE",
          7:"E6F1FB",8:"E6F1FB",9:"EEEDFE",10:"EEEDFE"}  # 等级配色带
    thin=Side(style="thin",color="DDDDDD"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
    ws.append(COLS)
    for c in range(1,len(COLS)+1):
        cell=ws.cell(1,c); cell.fill=head_fill; cell.font=head_font
        cell.alignment=Alignment(horizontal="center",vertical="center")
    for r in rows:
        ws.append(r)
        L=r[3]; fill=PatternFill("solid",fgColor=band[L])
        ws.cell(ws.max_row,4).fill=fill   # 给"阶级"列上色,直观标注 1-10 级
        ws.cell(ws.max_row,4).alignment=Alignment(horizontal="center")
    widths=[12,26,9,6,16,24,16,11,13,14,16,22,11,11,12,9,14,11,10,9,12,12]
    for i,wd in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=wd
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(COLS))}{ws.max_row}"
    for row in ws.iter_rows(min_row=1,max_row=ws.max_row,max_col=len(COLS)):
        for cell in row: cell.border=border

    # 覆盖矩阵
    mx=wb.create_sheet("覆盖矩阵")
    subs=[s[0] for s in SUBJECTS]
    mx.append(["学科"]+[f"{i}阶" for i in range(1,11)]+["合计"])
    for s in subs:
        row=[s]
        for L in range(1,11):
            row.append(f'=COUNTIFS(项目总表!C:C,"{s}",项目总表!D:D,{L})')
        row.append(f'=SUM(B{mx.max_row+1}:K{mx.max_row+1})')
        mx.append(row)
    tot=["合计"]+[f'=SUM({get_column_letter(c)}2:{get_column_letter(c)}{1+len(subs)})' for c in range(2,12)]
    tot.append(f'=SUM(L2:L{1+len(subs)})'); mx.append(tot)
    for c in range(1,13): mx.cell(1,c).font=Font(bold=True)
    wb.save(path); print("XLSX ->",path, len(rows),"rows")

if __name__=="__main__":
    rows=build_rows()
    assert len(rows)==600, len(rows)
    write_csv(rows)
    try:
        write_xlsx(rows)
    except Exception as e:
        print("xlsx 跳过(缺 openpyxl?):",e)
